import asyncio
import base64
import difflib
import logging
import os
import re
import shutil
import sys
import tempfile
from datetime import date as _date
from pathlib import Path

from claude_agent_sdk import (
    query, ClaudeAgentOptions, AssistantMessage, ResultMessage,
    ToolUseBlock, ToolResultBlock, ThinkingBlock, TextBlock, UserMessage,
)
from fastapi import APIRouter, HTTPException, Request
from pydantic import BaseModel

from lib.agent_env import agent_env
from lib.auth import verify_auth
from lib.job_progress import post_event
from lib.skill_loader import load_skill, SKILLS_DIR

logger = logging.getLogger(__name__)

router = APIRouter()

# Redis cancel-key polling — mirrors the Node worker's `chat-job:cancel:<jobId>` key.
# A module-level client is created lazily on first use.
_redis_client = None


def _get_redis_client():
    global _redis_client
    if _redis_client is None:
        import redis.asyncio as aioredis
        host = os.environ.get('REDIS_HOST', 'redis')
        port = int(os.environ.get('REDIS_PORT', '6379'))
        is_elasticache = host.endswith('.cache.amazonaws.com')
        password = os.environ.get('REDIS_PASSWORD', '')
        url = f'rediss://{host}:{port}' if is_elasticache else f'redis://{host}:{port}'
        _redis_client = aioredis.from_url(
            url,
            decode_responses=True,
            **(({'ssl_cert_reqs': None}) if is_elasticache else
               ({'password': password} if password else {})),
        )
    return _redis_client


async def _is_job_cancelled(job_id: str) -> bool:
    try:
        client = _get_redis_client()
        return await client.exists(f'chat-job:cancel:{job_id}') > 0
    except Exception:
        return False


def _start_cancel_watch(job_id: str) -> asyncio.Task:
    target = asyncio.current_task()
    async def _watch() -> None:
        while target and not target.done():
            await asyncio.sleep(2)
            if await _is_job_cancelled(job_id):
                if not target.done():
                    target.cancel()
                return
    return asyncio.create_task(_watch())


class CreateDocxRequest(BaseModel):
    model_id: str
    user_id: str
    chat_message_id: str | None = None
    user_group_id: str | None = None
    content: str
    title: str = ''
    job_id: str | None = None


class EditDocxRequest(BaseModel):
    model_id: str
    user_id: str
    chat_message_id: str | None = None
    user_group_id: str | None = None
    data_b64: str
    filename: str = ''
    changes: str
    job_id: str | None = None


class CreateHtmlRequest(BaseModel):
    model_id: str
    user_id: str
    chat_message_id: str | None = None
    user_group_id: str | None = None
    content: str
    title: str = ''
    style: str = 'default'
    feedback: str = ''
    job_id: str | None = None


class CreateMp4Request(BaseModel):
    model_id: str
    user_id: str
    chat_message_id: str | None = None
    user_group_id: str | None = None
    content: str
    title: str = ''
    style: str = 'default'
    feedback: str = ''
    job_id: str | None = None


@router.post('/create-docx')
async def create_docx(request: Request, body: CreateDocxRequest):
    verify_auth(request)

    output_dir = Path(tempfile.mkdtemp(prefix='docx-'))
    output_path = output_dir / 'document.docx'

    try:
        today = _date.today().strftime('%B %d, %Y')
        logger.info('[CLAUDE-SERVICE] create-docx title=%r', body.title[:60])

        skill_prompt = load_skill('document')
        prompt = (
            f'{skill_prompt}\n\n'
            f'---\n\n'
            f'## Task\n\n'
            f'Create a Word document from the following content.\n\n'
            f'Title: {body.title}\n'
            f'Today\'s date: {today}\n\n'
            f'Content:\n{body.content}\n\n'
            f'Write the complete file to: {output_path}\n'
            f'Print `DONE: {output_path}` when complete.\n'
        )

        options = ClaudeAgentOptions(
            tools=['Bash'],
            allowed_tools=['Bash'],
            model=body.model_id,
            env=agent_env(body.user_id, body.chat_message_id, body.user_group_id),
        )

        if body.job_id:
            await post_event(body.job_id, {'type': 'subagent_progress', 'label': 'Generating document...'})

        source_script: str | None = None
        gen_step = 0
        pending_tool_calls: dict[str, dict] = {}
        _cancel_watcher = _start_cancel_watch(body.job_id) if body.job_id else None
        try:
            async for message in query(prompt=prompt, options=options):
                if isinstance(message, AssistantMessage):
                    for block in message.content:
                        if isinstance(block, ToolUseBlock):
                            gen_step += 1
                            raw_input = block.input if isinstance(block.input, dict) else {}
                            pending_tool_calls[block.id] = {'tool': block.name, 'step': gen_step}
                            if block.name == 'Bash' and isinstance(raw_input, dict):
                                cmd = raw_input.get('command', '')
                                if "require('docx')" in cmd or 'Packer' in cmd:
                                    source_script = cmd
                            if body.job_id:
                                await post_event(body.job_id, {
                                    'type': 'subagent_tool_call',
                                    'tool': block.name,
                                    'args': _sanitize_tool_args(block.name, raw_input),
                                    'step': gen_step,
                                })
                        elif isinstance(block, ThinkingBlock) and body.job_id:
                            await post_event(body.job_id, {
                                'type': 'subagent_thinking',
                                'step': gen_step,
                                'length': len(block.thinking),
                            })
                elif isinstance(message, UserMessage) and body.job_id:
                    if isinstance(message.content, list):
                        for block in message.content:
                            if isinstance(block, ToolResultBlock):
                                call_info = pending_tool_calls.pop(block.tool_use_id, None)
                                content_len = len(block.content) if isinstance(block.content, str) else 0
                                await post_event(body.job_id, {
                                    'type': 'subagent_tool_result',
                                    'tool': call_info['tool'] if call_info else 'unknown',
                                    'step': call_info['step'] if call_info else gen_step,
                                    'contentLength': content_len,
                                    'isError': block.is_error or False,
                                })
                elif isinstance(message, ResultMessage) and body.job_id:
                    await post_event(body.job_id, {
                        'type': 'subagent_complete',
                        'label': 'Document generation complete',
                        'durationMs': message.duration_ms,
                    })
        except asyncio.CancelledError:
            logger.info('[CLAUDE-SERVICE] job cancelled, stopping create-docx')
        finally:
            if _cancel_watcher:
                _cancel_watcher.cancel()

        if not output_path.exists():
            raise HTTPException(status_code=422, detail='Document generation failed — agent did not produce a file')

        raw = output_path.read_bytes()
        content = base64.b64encode(raw).decode('ascii')

        section_count: int | None = None
        try:
            from docx import Document as _Document
            _doc = _Document(str(output_path))
            section_count = sum(1 for p in _doc.paragraphs if p.style.name.startswith('Heading 1'))
        except Exception:
            pass

        logger.info('[CLAUDE-SERVICE] create-docx complete sections=%s', section_count)
        return {
            'artifacts': {
                'content': content,
                'extension': '.docx',
                'encoding': 'base64',
            },
            'source_script': source_script,
            'section_count': section_count,
        }
    finally:
        shutil.rmtree(output_dir, ignore_errors=True)


@router.post('/edit-docx')
async def edit_docx_endpoint(request: Request, body: EditDocxRequest):
    verify_auth(request)

    output_dir = Path(tempfile.mkdtemp(prefix='docx-edit-'))
    output_path = output_dir / 'document.docx'

    try:
        today = _date.today().strftime('%B %d, %Y')
        logger.info('[CLAUDE-SERVICE] edit-docx changes=%r', body.changes[:120])

        skill_prompt = load_skill('document')

        if not body.data_b64:
            raise HTTPException(status_code=422, detail='edit-docx requires data_b64')

        raw_suffix = Path(body.filename).suffix.lower() if body.filename else '.docx'
        input_suffix = raw_suffix if raw_suffix in ('.docx', '.dotx') else '.docx'
        input_path = output_dir / f'input{input_suffix}'
        input_path.write_bytes(base64.b64decode(body.data_b64))
        unpack_dir = output_dir / 'unpacked'
        merge_runs_script = SKILLS_DIR / 'document' / 'merge_runs.py'
        prompt = (
            f'{skill_prompt}\n\n'
            f'---\n\n'
            f'## Task\n\n'
            f'Edit the Word document using the unzip/XML/rezip approach '
            f'(docx-js cannot open existing files).\n\n'
            f'Today\'s date: {today}\n'
            f'Input file: {input_path}\n'
            f'Unpack to: {unpack_dir}\n'
            f'Output file: {output_path}\n\n'
            f'Requested changes:\n{body.changes}\n\n'
            f'## Length reduction rule\n\n'
            f'If the requested changes ask to shorten, condense, reduce to a page count, or make '
            f'the document more concise, follow this order strictly:\n\n'
            f'1. **Rewrite the text to be shorter.** Cut sentences, tighten wording, remove '
            f'redundant sections. This is the only acceptable way to reduce length.\n'
            f'2. Do NOT reduce font sizes. Do NOT reduce spacing or margins. Do NOT squeeze '
            f'the layout. A document with tiny fonts or no breathing room looks broken and is '
            f'worse than a document that is slightly too long.\n'
            f'3. After rewriting, render once to check the page count. If it still overflows '
            f'by a small amount, cut a little more text — do not touch fonts or spacing.\n'
            f'4. It is acceptable to end up slightly over the target page count if the content '
            f'cannot be cut further without losing meaning. Never sacrifice readability for a '
            f'hard page count.\n\n'
            f'Steps:\n'
            f'1. `unzip -q {input_path} -d {unpack_dir}/`\n'
            f'2. `python {merge_runs_script} {unpack_dir}/` — coalesces fragmented runs so text is findable\n'
            f'3. Edit `{unpack_dir}/word/document.xml` in place — do NOT reformat or pretty-print.\n'
            f'4. `(cd {unpack_dir} && rm -f {output_path} && zip -Xr {output_path} .)`\n'
            f'5. Verify the output renders correctly with soffice + pdftoppm, fix if needed.\n'
            f'6. Print `DONE: {output_path}` when complete.\n'
        )

        options = ClaudeAgentOptions(
            tools=['Bash'],
            allowed_tools=['Bash'],
            model=body.model_id,
            env=agent_env(body.user_id, body.chat_message_id, body.user_group_id),
        )

        if body.job_id:
            await post_event(body.job_id, {'type': 'subagent_progress', 'label': 'Applying edits to document...'})

        source_script: str | None = None
        gen_step = 0
        pending_tool_calls: dict[str, dict] = {}
        _cancel_watcher = _start_cancel_watch(body.job_id) if body.job_id else None
        try:
            async for message in query(prompt=prompt, options=options):
                if isinstance(message, AssistantMessage):
                    for block in message.content:
                        if isinstance(block, ToolUseBlock):
                            gen_step += 1
                            raw_input = block.input if isinstance(block.input, dict) else {}
                            pending_tool_calls[block.id] = {'tool': block.name, 'step': gen_step}
                            if block.name == 'Bash' and isinstance(raw_input, dict):
                                cmd = raw_input.get('command', '')
                                if "require('docx')" in cmd or 'Packer' in cmd:
                                    source_script = cmd
                            if body.job_id:
                                await post_event(body.job_id, {
                                    'type': 'subagent_tool_call',
                                    'tool': block.name,
                                    'args': _sanitize_tool_args(block.name, raw_input),
                                    'step': gen_step,
                                })
                        elif isinstance(block, ThinkingBlock) and body.job_id:
                            await post_event(body.job_id, {
                                'type': 'subagent_thinking',
                                'step': gen_step,
                                'length': len(block.thinking),
                            })
                elif isinstance(message, UserMessage) and body.job_id:
                    if isinstance(message.content, list):
                        for block in message.content:
                            if isinstance(block, ToolResultBlock):
                                call_info = pending_tool_calls.pop(block.tool_use_id, None)
                                content_len = len(block.content) if isinstance(block.content, str) else 0
                                await post_event(body.job_id, {
                                    'type': 'subagent_tool_result',
                                    'tool': call_info['tool'] if call_info else 'unknown',
                                    'step': call_info['step'] if call_info else gen_step,
                                    'contentLength': content_len,
                                    'isError': block.is_error or False,
                                })
                elif isinstance(message, ResultMessage) and body.job_id:
                    await post_event(body.job_id, {
                        'type': 'subagent_complete',
                        'label': 'Document edit complete',
                        'durationMs': message.duration_ms,
                    })
        except asyncio.CancelledError:
            logger.info('[CLAUDE-SERVICE] job cancelled, stopping edit-docx')
        finally:
            if _cancel_watcher:
                _cancel_watcher.cancel()

        if not output_path.exists():
            raise HTTPException(status_code=422, detail='Document edit failed — agent did not produce a file')

        raw = output_path.read_bytes()
        content = base64.b64encode(raw).decode('ascii')

        section_count: int | None = None
        diff_stat: dict | None = None
        try:
            from docx import Document as _Document
            before_text = '\n'.join(p.text for p in _Document(str(input_path)).paragraphs)
            _doc = _Document(str(output_path))
            after_text = '\n'.join(p.text for p in _doc.paragraphs)
            section_count = sum(1 for p in _doc.paragraphs if p.style.name.startswith('Heading 1'))
            diff_stat = _line_diff_stat(before_text, after_text)
        except Exception:
            pass

        logger.info('[CLAUDE-SERVICE] edit-docx complete sections=%s', section_count)
        return {
            'artifacts': {
                'content': content,
                'extension': '.docx',
                'encoding': 'base64',
            },
            'source_script': source_script,
            'section_count': section_count,
            'diffStat': diff_stat,
        }
    finally:
        shutil.rmtree(output_dir, ignore_errors=True)


@router.post('/create-html')
async def create_html(request: Request, body: CreateHtmlRequest):
    verify_auth(request)

    job_id = body.job_id or os.urandom(8).hex()
    output_dir = Path(tempfile.mkdtemp(prefix='html-'))
    output_path = output_dir / 'html-presentation.html'

    try:
        logger.info('[CLAUDE-SERVICE] create-html title=%r', body.title[:60])

        result = await _run_agent_skill(
            skill_name='html-presentation',
            output_type='HTML presentation',
            output_path=output_path,
            model_id=body.model_id,
            user_id=body.user_id,
            chat_message_id=body.chat_message_id,
            user_group_id=body.user_group_id,
            title=body.title,
            content=body.content,
            style=body.style,
            feedback=body.feedback,
            job_id=body.job_id,
        )

        if not output_path.exists():
            raise HTTPException(status_code=422, detail='HTML presentation generation failed')

        content = output_path.read_text(encoding='utf-8')

        return {'artifacts': {
            'content': content,
            'extension': '.html',
            'encoding': 'utf-8',
        }}
    finally:
        shutil.rmtree(output_dir, ignore_errors=True)


class CreatePptxRequest(BaseModel):
    model_id: str
    internal_model_id: str = ''
    user_id: str
    chat_message_id: str | None = None
    user_group_id: str | None = None
    content: str
    title: str = ''
    use_template: bool = True
    job_id: str | None = None


class CreateXlsxRequest(BaseModel):
    model_id: str
    user_id: str
    chat_message_id: str | None = None
    user_group_id: str | None = None
    content: str
    title: str = ''
    job_id: str | None = None


class EditXlsxRequest(BaseModel):
    model_id: str
    user_id: str
    chat_message_id: str | None = None
    user_group_id: str | None = None
    source_script: str = ''
    data_b64: str = ''
    filename: str = ''
    changes: str
    title: str = ''
    job_id: str | None = None


async def _resolve_pptx_template(user_id: str, env: dict) -> bytes | None:
    """Fetch group template via /v1/templates/resolve. Returns bytes or None."""
    import httpx as _httpx
    try:
        async with _httpx.AsyncClient(timeout=30) as client:
            resp = await client.post(
                f'{env["ANTHROPIC_BASE_URL"]}/v1/templates/resolve',
                headers={'Authorization': f'Bearer {env["ANTHROPIC_API_KEY"]}'},
                json={'userId': user_id, 'fileExtension': 'pptx'},
            )
            if resp.status_code == 404:
                return None
            resp.raise_for_status()
            return base64.b64decode(resp.json()['fileData'])
    except Exception as exc:
        logger.warning('[CLAUDE-SERVICE] template resolve failed: %s %s', type(exc).__name__, exc)
        return None




def _is_pptx_script(cmd: str) -> bool:
    return ('.save(' in cmd or 'from pptx' in cmd) and len(cmd) > 200



def _fix_potx_content_type(template_path: Path) -> None:
    """Convert .potx content type to .pptx so python-pptx accepts it."""
    import zipfile as _zf
    import tempfile as _tf

    with _zf.ZipFile(template_path, 'r') as zin:
        ct_xml = zin.read('[Content_Types].xml').decode('utf-8')
        if 'presentationml.template' not in ct_xml:
            return

        ct_xml = ct_xml.replace(
            'application/vnd.openxmlformats-officedocument.presentationml.template.main+xml',
            'application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml',
        )

        tmp_fd, tmp_path = _tf.mkstemp(suffix='.pptx')
        os.close(tmp_fd)
        with _zf.ZipFile(tmp_path, 'w', _zf.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                if item.filename == '[Content_Types].xml':
                    zout.writestr(item, ct_xml.encode('utf-8'))
                else:
                    zout.writestr(item, zin.read(item.filename))

    Path(tmp_path).replace(template_path)


def _strip_template_slides(template_path: Path) -> None:
    """Remove all existing slides from the template, keeping only layouts/masters."""
    try:
        from pptx import Presentation as _Presentation
        prs = _Presentation(str(template_path))
        while len(prs.slides) > 0:
            rId = prs.slides._sldIdLst[0].get(
                '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id'
            )
            prs.part.drop_rel(rId)
            prs.slides._sldIdLst.remove(prs.slides._sldIdLst[0])
        prs.save(str(template_path))
    except Exception as exc:
        logger.warning('[CLAUDE-SERVICE] strip template slides failed: %s', exc)


def _discover_template_layouts(template_path: Path) -> dict:
    """Discover layouts, returning a dict with layout info and access instructions.

    Some templates have multiple slide masters:
    - Master 0 (slideMaster1): Title layouts with large logo group
    - Master 1 (slideMaster2): Content layouts with small logo + footer

    Returns dict with keys: 'title_layout_access', 'content_layout_access', 'layout_info', 'content_master_idx'
    """
    try:
        _fix_potx_content_type(template_path)
        _strip_template_slides(template_path)
        from pptx import Presentation as _Presentation
        prs = _Presentation(str(template_path))

        result = {
            'title_layout_access': None,
            'content_layout_access': None,
            'layout_info': '',
            'content_master_idx': None,
        }

        lines = []

        # Find the content master (has "Content - Blank" or similar content layouts)
        content_master_idx = None
        for mi, master in enumerate(prs.slide_masters):
            for layout in master.slide_layouts:
                if 'Content' in layout.name and 'Blank' in layout.name:
                    content_master_idx = mi
                    break
            if content_master_idx is not None:
                break

        # Title layouts (from top-level prs.slide_layouts)
        lines.append('**Title layouts** (via `prs.slide_layouts[i]`):')
        for i, layout in enumerate(prs.slide_layouts):
            if 'Title' not in layout.name and 'Divider' not in layout.name:
                continue
            placeholders = []
            try:
                for ph in layout.placeholders:
                    placeholders.append(f'idx={ph.placeholder_format.idx} "{ph.name}"')
            except Exception:
                pass
            ph_str = ', '.join(placeholders) if placeholders else 'no placeholders'
            lines.append(f'- `prs.slide_layouts[{i}]`: **{layout.name}** — {ph_str}')
            if result['title_layout_access'] is None and 'Color' in layout.name:
                result['title_layout_access'] = f'prs.slide_layouts[{i}]'

        # Content layouts (from content master)
        if content_master_idx is not None:
            result['content_master_idx'] = content_master_idx
            content_master = prs.slide_masters[content_master_idx]
            lines.append(f'\n**Content layouts** (via `prs.slide_masters[{content_master_idx}].slide_layouts[i]`):')
            for i, layout in enumerate(content_master.slide_layouts):
                placeholders = []
                try:
                    for ph in layout.placeholders:
                        placeholders.append(f'idx={ph.placeholder_format.idx} "{ph.name}"')
                except Exception:
                    pass
                ph_str = ', '.join(placeholders) if placeholders else 'no placeholders'
                lines.append(f'- `prs.slide_masters[{content_master_idx}].slide_layouts[{i}]`: **{layout.name}** — {ph_str}')
                if result['content_layout_access'] is None and 'Blank' in layout.name:
                    result['content_layout_access'] = f'prs.slide_masters[{content_master_idx}].slide_layouts[{i}]'

            # Describe what the content master provides automatically
            master_shapes = []
            for shape in content_master.shapes:
                try:
                    if shape.placeholder_format is not None:
                        continue
                except Exception:
                    pass
                left = shape.left / 914400 if shape.left else 0
                top = shape.top / 914400 if shape.top else 0
                width = shape.width / 914400 if shape.width else 0
                height = shape.height / 914400 if shape.height else 0
                master_shapes.append(f'{shape.name} at ({left:.2f}", {top:.2f}") size ({width:.2f}" x {height:.2f}")')
            if master_shapes:
                lines.append(f'\n**Content master provides automatically** (do NOT add these manually):')
                for ms in master_shapes:
                    lines.append(f'- {ms}')

        result['layout_info'] = '\n'.join(lines)
        return result
    except Exception as exc:
        return {'title_layout_access': None, 'content_layout_access': None,
                'layout_info': f'(Could not discover layouts: {exc})', 'content_master_idx': None}


def _extract_template_logo(template_path: Path, output_dir: Path) -> Path | None:
    """Extract the largest PNG from the template (likely the logo). Returns path or None."""
    import zipfile as _zf

    try:
        with _zf.ZipFile(template_path) as zf:
            pngs = [(n, zf.getinfo(n).file_size) for n in zf.namelist()
                    if n.startswith('ppt/media/') and n.lower().endswith('.png')]
            if not pngs:
                return None

            # The logo is typically the largest PNG in the first layout
            pngs.sort(key=lambda x: x[1], reverse=True)
            logo_name = pngs[0][0]

            logo_path = output_dir / 'logo.png'
            logo_path.write_bytes(zf.read(logo_name))
            return logo_path
    except Exception:
        return None


def _extract_template_text_sizes(template_path: Path) -> dict:
    """Extract text sizes from the template's slide master txStyles.

    Reads titleStyle/bodyStyle from the first two masters:
    - Master 1 (title master): title size, body=subtitle size
    - Master 2 (content master): title=header size, body=content body size

    Returns dict with keys: 'title_pt', 'subtitle_pt', 'date_pt', 'body_pt', 'header_pt'.
    """
    defaults = {'title_pt': 36, 'subtitle_pt': 20, 'date_pt': 12, 'body_pt': 12, 'header_pt': 24}
    try:
        import zipfile
        from xml.etree import ElementTree as ET

        ns_a = 'http://schemas.openxmlformats.org/drawingml/2006/main'
        ns_p = 'http://schemas.openxmlformats.org/presentationml/2006/main'

        sizes = {}

        with zipfile.ZipFile(template_path) as zf:
            masters = sorted([n for n in zf.namelist()
                              if 'slideMaster' in n and n.endswith('.xml') and '/_rels/' not in n])

            for idx, master_file in enumerate(masters[:2]):
                root = ET.fromstring(zf.read(master_file))

                for tx_styles in root.iter(f'{{{ns_p}}}txStyles'):
                    # titleStyle → lvl1pPr defRPr sz
                    for title_style in tx_styles.iter(f'{{{ns_p}}}titleStyle'):
                        for lvl in title_style:
                            def_rpr = lvl.find(f'{{{ns_a}}}defRPr')
                            if def_rpr is not None:
                                sz = def_rpr.get('sz')
                                if sz:
                                    pt = int(sz) // 100
                                    if idx == 0:
                                        sizes['title_pt'] = pt
                                    else:
                                        sizes['header_pt'] = pt
                                break

                    # bodyStyle → lvl1pPr defRPr sz
                    for body_style in tx_styles.iter(f'{{{ns_p}}}bodyStyle'):
                        for lvl in body_style:
                            def_rpr = lvl.find(f'{{{ns_a}}}defRPr')
                            if def_rpr is not None:
                                sz = def_rpr.get('sz')
                                if sz:
                                    pt = int(sz) // 100
                                    if idx == 0:
                                        sizes['subtitle_pt'] = pt
                                    else:
                                        sizes['body_pt'] = pt
                                break

        # Date is typically 60-65% of subtitle size
        if 'subtitle_pt' in sizes and 'date_pt' not in sizes:
            sizes['date_pt'] = max(10, round(sizes['subtitle_pt'] * 0.6))

        return {**defaults, **sizes}
    except Exception:
        return defaults


def _extract_template_geometry(template_path: Path, title_layout_access: str, content_layout_access: str) -> dict:
    """Extract placeholder positions and content bounds from the template.

    Returns dict with:
    - title_boxes: list of {x, y, w, h} for title layout placeholders (title, subtitle, date)
    - content_top: y where content starts (below logo/header)
    - content_bottom: y where content ends (above footer)
    - content_left: x left margin
    - content_right: x right margin
    - logo_w, logo_h: logo dimensions
    - footer_y: footer y position
    """
    defaults = {
        'title_boxes': [
            {'x': 0.5, 'y': 2.0, 'w': 5.0, 'h': 1.5, 'role': 'title'},
            {'x': 0.5, 'y': 3.7, 'w': 5.0, 'h': 1.0, 'role': 'subtitle'},
            {'x': 0.5, 'y': 4.8, 'w': 5.0, 'h': 0.4, 'role': 'date'},
        ],
        'content_top': 0.9,
        'content_bottom': 6.75,
        'content_left': 0.5,
        'content_right': 12.8,
        'logo_w': 0.83,
        'logo_h': 0.71,
        'footer_y': 6.97,
    }
    try:
        from pptx import Presentation as _Presentation
        prs = _Presentation(str(template_path))

        # Parse the layout access strings to get actual layout objects
        title_layout = None
        content_master = None

        # Evaluate title layout access (e.g. "prs.slide_layouts[3]")
        if 'slide_masters' in title_layout_access:
            parts = title_layout_access.replace('prs.', '').split('.')
            obj = prs
            for part in parts:
                if '[' in part:
                    name = part[:part.index('[')]
                    idx = int(part[part.index('[') + 1:part.index(']')])
                    obj = getattr(obj, name)[idx]
                else:
                    obj = getattr(obj, part)
            title_layout = obj
        elif 'slide_layouts' in title_layout_access:
            idx = int(title_layout_access.split('[')[1].split(']')[0])
            title_layout = prs.slide_layouts[idx]

        # Get content master
        if 'slide_masters' in content_layout_access:
            master_idx = int(content_layout_access.split('slide_masters[')[1].split(']')[0])
            content_master = prs.slide_masters[master_idx]

        result = {}

        # Extract title layout placeholder positions
        if title_layout:
            emu_to_inches = 914400
            boxes = []
            ph_map = {}
            for ph in title_layout.placeholders:
                ph_type = ph.placeholder_format.type
                idx = ph.placeholder_format.idx
                x = round(ph.left / emu_to_inches, 2) if ph.left else 0
                y = round(ph.top / emu_to_inches, 2) if ph.top else 0
                w = round(ph.width / emu_to_inches, 2) if ph.width else 0
                h = round(ph.height / emu_to_inches, 2) if ph.height else 0

                from pptx.enum.shapes import PP_PLACEHOLDER
                if ph_type == PP_PLACEHOLDER.CENTER_TITLE or ph_type == PP_PLACEHOLDER.TITLE:
                    boxes.append({'x': x, 'y': y, 'w': w, 'h': h, 'role': 'title'})
                elif ph_type == PP_PLACEHOLDER.SUBTITLE:
                    boxes.append({'x': x, 'y': y, 'w': w, 'h': h, 'role': 'subtitle'})
                elif ph_type == PP_PLACEHOLDER.DATE:
                    boxes.append({'x': x, 'y': y, 'w': w, 'h': h, 'role': 'date'})
                elif ph_type == PP_PLACEHOLDER.BODY:
                    boxes.append({'x': x, 'y': y, 'w': w, 'h': h, 'role': 'subtitle'})

            if boxes:
                result['title_boxes'] = sorted(boxes, key=lambda b: b['y'])

        # Extract content area bounds from content master shapes
        if content_master:
            emu_to_inches = 914400
            logo_h = 0.0
            footer_y = 7.5
            logo_w = 0.0

            for shape in content_master.shapes:
                try:
                    if shape.placeholder_format is not None:
                        continue
                except Exception:
                    pass
                top = round(shape.top / emu_to_inches, 2) if shape.top else 0
                left = round(shape.left / emu_to_inches, 2) if shape.left else 0
                height = round(shape.height / emu_to_inches, 2) if shape.height else 0
                width = round(shape.width / emu_to_inches, 2) if shape.width else 0

                # Logo is typically at top-left, small
                if top < 1.0 and left < 2.0 and height < 1.5:
                    logo_h = max(logo_h, height)
                    logo_w = max(logo_w, width)

                # Footer is typically near bottom
                if top > 6.0:
                    footer_y = min(footer_y, top)

            if logo_h > 0:
                result['logo_h'] = logo_h
                result['logo_w'] = logo_w
                result['content_top'] = round(logo_h + 0.2, 2)

            if footer_y < 7.5:
                result['footer_y'] = round(footer_y, 2)
                result['content_bottom'] = round(footer_y - 0.22, 2)

            # Content margins: 0.5" from edges
            result['content_left'] = 0.5
            result['content_right'] = round(prs.slide_width / emu_to_inches - 0.5, 2)

        return {**defaults, **result}
    except Exception:
        return defaults


def _extract_template_brand(template_path: Path) -> dict:
    """Extract brand colors and fonts from a .pptx template's theme XML.

    Returns dict with keys: 'text' (prompt string), 'colors' (dict), 'fonts' (dict).
    """
    import zipfile
    from xml.etree import ElementTree as ET

    empty = {'text': '', 'colors': {}, 'fonts': {}}
    try:
        with zipfile.ZipFile(template_path) as zf:
            theme_files = [n for n in zf.namelist() if 'theme' in n.lower() and n.endswith('.xml')]
            if not theme_files:
                return empty

            root = ET.fromstring(zf.read(theme_files[0]))
            ns_a = 'http://schemas.openxmlformats.org/drawingml/2006/main'

            colors = {}
            for scheme in root.iter(f'{{{ns_a}}}clrScheme'):
                for child in scheme:
                    tag = child.tag.split('}')[-1]
                    for attr_elem in child:
                        val = attr_elem.get('val') or attr_elem.get('lastClr')
                        if val:
                            colors[tag] = val
                            break

            fonts = {}
            for font_scheme in root.iter(f'{{{ns_a}}}fontScheme'):
                for major in font_scheme.iter(f'{{{ns_a}}}majorFont'):
                    for latin in major.iter(f'{{{ns_a}}}latin'):
                        fonts['heading'] = latin.get('typeface')
                for minor in font_scheme.iter(f'{{{ns_a}}}minorFont'):
                    for latin in minor.iter(f'{{{ns_a}}}latin'):
                        fonts['body'] = latin.get('typeface')

            lines = ['### Brand Theme (from uploaded template)\n']
            if colors:
                lines.append('**Colors (use these — do NOT invent colors):**')
                for name, hex_val in colors.items():
                    lines.append(f'- {name}: `#{hex_val}` → `RGBColor(0x{hex_val[0:2]}, 0x{hex_val[2:4]}, 0x{hex_val[4:6]})`')
                lines.append('')
            if fonts:
                lines.append('**Fonts (set on every text run):**')
                for role, face in fonts.items():
                    lines.append(f'- {role}: `{face}` → `font.name = "{face}"`')
                lines.append('')

            lines.append('Use ONLY these colors. accent5 = primary accent (header bars). '
                         'accent1 = secondary (highlights). dk1 = text. lt1 = background/white text on dark.')
            return {'text': '\n'.join(lines), 'colors': colors, 'fonts': fonts}
    except Exception:
        return empty


@router.post('/create-pptx')
async def create_pptx(request: Request, body: CreatePptxRequest):
    verify_auth(request)

    output_dir = Path(tempfile.mkdtemp(prefix='pptx-'))
    output_path = output_dir / 'presentation.pptx'

    try:
        logger.info('[CLAUDE-SERVICE] create-pptx title=%r', body.title[:60])

        source_script, gen_step, pending_tool_calls = '', 0, {}

        skill_prompt = load_skill('pptx')
        env = agent_env(body.user_id, body.chat_message_id, body.user_group_id)

        template_bytes = await _resolve_pptx_template(body.user_id, env) if body.use_template else None

        if template_bytes:
            template_path = output_dir / 'template.pptx'
            template_path.write_bytes(template_bytes)

            if body.job_id:
                await post_event(body.job_id, {'type': 'subagent_progress', 'label': 'Preparing branded template...'})

            # Discover available layouts and their placeholder info
            layout_result = _discover_template_layouts(template_path)
            layout_info = layout_result['layout_info']
            content_layout_access = layout_result['content_layout_access'] or 'prs.slide_layouts[1]'
            title_layout_access = layout_result['title_layout_access'] or 'prs.slide_layouts[0]'
            logger.info('[CLAUDE-SERVICE] template layouts: %s', layout_info[:500])
            logger.info('[CLAUDE-SERVICE] content_layout=%s title_layout=%s', content_layout_access, title_layout_access)

            brand = _extract_template_brand(template_path)
            brand_text = brand['text']
            brand_colors = brand['colors']
            brand_fonts = brand['fonts']
            font_name = brand_fonts.get('heading') or brand_fonts.get('body') or 'Inter'
            accent_primary = brand_colors.get('accent5', '00838F')
            accent_secondary = brand_colors.get('accent1', '23D2D7')
            logo_path = _extract_template_logo(template_path, output_dir)
            text_sizes = _extract_template_text_sizes(template_path)
            title_pt = text_sizes['title_pt']
            subtitle_pt = text_sizes['subtitle_pt']
            date_pt = text_sizes['date_pt']
            body_pt = text_sizes['body_pt']
            header_pt = text_sizes['header_pt']
            logger.info('[CLAUDE-SERVICE] template text sizes: %s', text_sizes)

            geo = _extract_template_geometry(template_path, title_layout_access, content_layout_access)
            title_boxes = geo['title_boxes']
            content_top = geo['content_top']
            content_bottom = geo['content_bottom']
            content_left = geo['content_left']
            content_right = geo['content_right']
            content_width = round(content_right - content_left, 2)
            content_height = round(content_bottom - content_top, 2)
            logo_w = geo['logo_w']
            logo_h = geo['logo_h']
            footer_y = geo['footer_y']
            # Map title boxes by role
            title_box = next((b for b in title_boxes if b['role'] == 'title'), title_boxes[0] if title_boxes else {'x': 0.5, 'y': 2.0, 'w': 5.0, 'h': 1.5})
            subtitle_box = next((b for b in title_boxes if b['role'] == 'subtitle'), title_boxes[1] if len(title_boxes) > 1 else {'x': 0.5, 'y': 3.7, 'w': 5.0, 'h': 1.0})
            date_box = next((b for b in title_boxes if b['role'] == 'date'), title_boxes[2] if len(title_boxes) > 2 else {'x': 0.5, 'y': 4.8, 'w': 5.0, 'h': 0.4})
            logger.info('[CLAUDE-SERVICE] template geometry: content_top=%s content_bottom=%s footer_y=%s title_boxes=%s', content_top, content_bottom, footer_y, title_boxes)

            template_section = (
                f'## Approach — BRANDED TEMPLATE (python-pptx)\n\n'
                f'A branded template is at: `{template_path}`\n\n'
                f'Use **python-pptx** to open this template and create slides.\n\n'
                f'{brand_text}\n\n'
                f'**Available layouts:**\n{layout_info}\n\n'
                f'## CRITICAL: TWO DIFFERENT LAYOUT SOURCES\n\n'
                f'This template has SEPARATE masters for title vs content slides:\n\n'
                f'- **Title/Closing slides**: `{title_layout_access}` — branded full-bleed background\n'
                f'- **Content slides**: `{content_layout_access}` — clean canvas with small logo + footer from master\n\n'
                f'The content master automatically provides:\n'
                f'- Small logo ({logo_w}" x {logo_h}") at top-left\n'
                f'- Footer text at bottom-left\n'
                f'- "Copyright" text at bottom-right\n'
                f'- Slide number at bottom-right\n\n'
                f'Do NOT add these elements manually — they come from the master!\n\n'
                f'## PLACEHOLDER REMOVAL (MANDATORY on EVERY slide)\n\n'
                f'Remove ALL placeholders immediately after adding any slide:\n'
                f'```python\n'
                f'from pptx.oxml.ns import qn\n'
                f'def strip_placeholders(slide):\n'
                f'    spTree = slide.shapes._spTree\n'
                f'    for sp in list(spTree):\n'
                f'        ph = sp.find(qn("p:nvSpPr") + "/" + qn("p:nvPr") + "/" + qn("p:ph"))\n'
                f'        if ph is not None:\n'
                f'            spTree.remove(sp)\n'
                f'```\n\n'
                f'## CRITICAL: USE set_text() FOR ALL TEXT — NEVER use tf.text or p.text directly\n\n'
                f'Because we strip placeholders, text has NO inherited styles. You MUST use this helper for ALL text:\n'
                f'```python\n'
                f'def set_text(tf, text, size, bold=False, color=RGBColor(0x00,0x00,0x00), font_name="{font_name}"):\n'
                f'    """Set text with explicit formatting. Use this for EVERY text element."""\n'
                f'    p = tf.paragraphs[0]\n'
                f'    p.text = text\n'
                f'    p.font.size = size\n'
                f'    p.font.bold = bold\n'
                f'    p.font.name = font_name\n'
                f'    p.font.color.rgb = color\n'
                f'```\n'
                f'**NEVER write `tf.text = "..."` or `p.text = "..."` without calling set_text or manually setting all font properties.**\n'
                f'Text without explicit size/font will render invisible or wrong. This is the #1 cause of broken decks.\n\n'
                f'## BREADCRUMB (REQUIRED on every content slide)\n\n'
                f'Every content slide MUST have a breadcrumb text line at the top, to the right of the logo.\n'
                f'This shows the section hierarchy (e.g. "Overview", "Architecture > Components").\n\n'
                f'```python\n'
                f'bc = slide.shapes.add_textbox(Inches({round(logo_w + 0.17, 2)}), Inches(0.15), Inches(5.5), Inches(0.4))\n'
                f'set_text(bc.text_frame, "Section Name", Pt({date_pt}), color=RGBColor(0x66,0x66,0x66))\n'
                f'```\n\n'
                f'Place it on EVERY content slide (not title/closing slides). Use the slide\'s section name.\n\n'
                f'## SAFE CONTENT AREA\n\n'
                f'**Content slides** (using the content master layout):\n'
                f'- **Top**: y = {content_top}" (below the {logo_h}" logo + whitespace)\n'
                f'- **Bottom**: y = {content_bottom}" (above footer text at {footer_y}")\n'
                f'- **Left**: x = {content_left}"\n'
                f'- **Right**: x = {content_right}"\n'
                f'- **Available height**: ~{content_height}" (from {content_top}" to {content_bottom}")\n'
                f'- **Available width**: ~{content_width}"\n\n'
                f'**Title slides** (Color layout — full-bleed branded background):\n'
                f'- EXACTLY 3 text boxes at x={title_box["x"]}" — title, subtitle, date. NO others.\n'
                f'- Do NOT add: taglines, italic text, "Prepared by" lines, org/division lines, or any extra elements\n'
                f'- Title: ({title_box["x"]}", {title_box["y"]}") {title_box["w"]}" x {title_box["h"]}" — {title_pt}pt bold black {font_name}\n'
                f'- Subtitle: ({subtitle_box["x"]}", {subtitle_box["y"]}") {subtitle_box["w"]}" x {subtitle_box["h"]}" — {subtitle_pt}pt regular black {font_name}\n'
                f'- Date: ({date_box["x"]}", {date_box["y"]}") {date_box["w"]}" x {date_box["h"]}" — {date_pt}pt bold black {font_name}\n\n'
                f'## CONTENT DENSITY — DO NOT OVERLOAD SLIDES\n\n'
                f'- Max 4-5 bullet points per slide\n'
                f'- Max 4 cards in a row\n'
                f'- Tables: max 5 rows per slide (split into multiple slides if more)\n'
                f'- Always 0.3" breathing room between elements\n'
                f'- **HARD LIMIT: No content below y={round(footer_y - 0.47, 2)}"** — the footer lives at {footer_y}". Anything below {round(footer_y - 0.47, 2)}" overlaps it.\n'
                f'- If content would overflow, SPLIT across multiple slides — NEVER cover the footer\n'
                f'- Better to have MORE slides than cramped slides\n\n'
                f'## ALIGNMENT — NUMBERED STEPS / ICON + TEXT PATTERNS\n\n'
                f'When placing a circle/icon next to text (e.g. numbered steps):\n'
                f'- The TEXT y-position must vertically center on the circle, not start at the circle\'s top\n'
                f'- Formula: text_y = circle_y + (circle_height - text_height) / 2\n'
                f'- Example: circle at y=2.0" h=0.6" → text_y = 2.0 + (0.6 - 0.4) / 2 = 2.1"\n'
                f'- Text should start AFTER the circle (x = circle_x + circle_width + 0.2" gap)\n'
                f'- NEVER let text overlap the circle horizontally\n\n'
                f'## TYPOGRAPHY — EXACT SIZES (match template precisely)\n\n'
                f'Font: {font_name} (set `font.name = "{font_name}"` on EVERY text run). These sizes are mandatory:\n\n'
                f'**Title slide (EXACTLY 3 text boxes — no italic, no extra lines):**\n'
                f'- Main title: Pt({title_pt}), bold, black\n'
                f'- Subtitle: Pt({subtitle_pt}), regular (NOT bold, NOT italic), black\n'
                f'- Date/org line: Pt({date_pt}), bold, black — format: "Org | Division  Month Year"\n\n'
                f'**Content slides:**\n'
                f'- Slide title (in accent5 header bar): Pt({header_pt}), bold, white\n'
                f'- Body text: Pt({body_pt}), regular\n'
                f'- Section subheadings: Pt({body_pt}), bold\n'
                f'- Card titles: Pt({body_pt}), bold\n'
                f'- Card body: Pt({body_pt - 2}), regular\n'
                f'- Breadcrumb: Pt({date_pt}), regular, gray\n'
                f'- Captions/footnotes: Pt({date_pt})\n\n'
                f'**Spacing rules:**\n'
                f'- After header bar: 0.15" gap before body content\n'
                f'- Between cards: 0.2" horizontal, 0.15" vertical\n'
                f'- Between bullet items: use paraSpaceAfter=Pt({max(3, date_pt // 3)})\n'
                f'- Between sections: 0.3" vertical gap\n\n'
                f'Do NOT deviate from these sizes — they are extracted from the template.\n\n'
                f'## Content Slide Styling\n\n'
                f'Style every content slide using the brand colors from the theme above (accent1-6):\n\n'
                f'- Add a **header bar** at y={content_top}" using the accent5 color, with white text (Pt({header_pt}) bold)\n'
                f'- Use accent5 for primary accent (header bars, key shapes)\n'
                f'- Use accent1 for secondary accent (lighter highlights, card borders)\n'
                f'- Use accent6 for tertiary elements if needed\n'
                f'- Reference the "Brand Theme" section above for exact hex values\n'
                f'- Add **icons** rendered via react-icons (example below)\n'
                f'- Add colored shape containers for card layouts\n'
                f'- NEVER leave a slide as just plain text on white\n\n'
                f'**Icon pattern:**\n'
                f'```python\n'
                f'import subprocess, os\n'
                f'icon_script = """\n'
                f"const React = require('react');\n"
                f"const ReactDOMServer = require('react-dom/server');\n"
                f"const {{ GiCookingPot }} = require('react-icons/gi');\n"
                f"const sharp = require('sharp');\n"
                f"const svg = ReactDOMServer.renderToStaticMarkup(React.createElement(GiCookingPot));\n"
                f"sharp(Buffer.from(svg)).resize(256).png().toBuffer().then(b => process.stdout.write(b));\n"
                f'"""\n'
                f'result = subprocess.run(["node", "-e", icon_script], capture_output=True,\n'
                f'    env={{**os.environ, "NODE_PATH": "/usr/local/lib/node_modules"}})\n'
                f'with open("/tmp/icon.png", "wb") as f: f.write(result.stdout)\n'
                f'slide.shapes.add_picture("/tmp/icon.png", Inches(1), Inches(2.5), Inches(0.5), Inches(0.5))\n'
                f'```\n\n'
                + f'**Logo**: The content master already shows the logo — do NOT add it manually.\n\n'
                + f'**Example (full working script):**\n'
                f'```python\n'
                f'from pptx import Presentation\n'
                f'from pptx.util import Inches, Pt\n'
                f'from pptx.dml.color import RGBColor\n'
                f'from pptx.enum.shapes import MSO_SHAPE\n'
                f'from pptx.enum.text import PP_ALIGN\n'
                f'from pptx.oxml.ns import qn\n\n'
                f'prs = Presentation("{template_path}")\n\n'
                f'# Access layouts from different masters\n'
                f'title_layout = {title_layout_access}   # branded full-bleed title\n'
                f'content_layout = {content_layout_access}  # clean + small logo from master\n\n'
                f'def strip_placeholders(slide):\n'
                f'    spTree = slide.shapes._spTree\n'
                f'    for sp in list(spTree):\n'
                f'        ph = sp.find(qn("p:nvSpPr") + "/" + qn("p:nvPr") + "/" + qn("p:ph"))\n'
                f'        if ph is not None:\n'
                f'            spTree.remove(sp)\n\n'
                f'def set_text(tf, text, size, bold=False, color=RGBColor(0x00,0x00,0x00)):\n'
                f'    """REQUIRED for all text. Never use tf.text directly."""\n'
                f'    p = tf.paragraphs[0]\n'
                f'    p.text = text\n'
                f'    p.font.size = size\n'
                f'    p.font.bold = bold\n'
                f'    p.font.name = "{font_name}"\n'
                f'    p.font.color.rgb = color\n\n'
                f'WHITE = RGBColor(0xFF,0xFF,0xFF)\n'
                f'BLACK = RGBColor(0x00,0x00,0x00)\n'
                f'GRAY = RGBColor(0x66,0x66,0x66)\n'
                f'DK2 = RGBColor(0x46,0x46,0x46)\n'
                f'ACCENT = RGBColor(0x{accent_primary[0:2]},0x{accent_primary[2:4]},0x{accent_primary[4:6]})\n\n'
                f'# --- Title slide (EXACTLY 3 text boxes) ---\n'
                f'slide = prs.slides.add_slide(title_layout)\n'
                f'strip_placeholders(slide)\n'
                f'tf = slide.shapes.add_textbox(Inches({title_box["x"]}), Inches({title_box["y"]}), Inches({title_box["w"]}), Inches({title_box["h"]})).text_frame\n'
                f'tf.word_wrap = True\n'
                f'set_text(tf, "Presentation Title", Pt({title_pt}), bold=True)\n'
                f'tf = slide.shapes.add_textbox(Inches({subtitle_box["x"]}), Inches({subtitle_box["y"]}), Inches({subtitle_box["w"]}), Inches({subtitle_box["h"]})).text_frame\n'
                f'tf.word_wrap = True\n'
                f'set_text(tf, "Subtitle goes here", Pt({subtitle_pt}))\n'
                f'tf = slide.shapes.add_textbox(Inches({date_box["x"]}), Inches({date_box["y"]}), Inches({date_box["w"]}), Inches({date_box["h"]})).text_frame\n'
                f'set_text(tf, "Organization  |  Month Year", Pt({date_pt}), bold=True)\n\n'
                f'# --- Content slide ---\n'
                f'slide = prs.slides.add_slide(content_layout)\n'
                f'strip_placeholders(slide)\n'
                f'# Breadcrumb\n'
                f'tf = slide.shapes.add_textbox(Inches({round(logo_w + 0.17, 2)}), Inches(0.15), Inches(5.5), Inches(0.4)).text_frame\n'
                f'set_text(tf, "Overview", Pt({date_pt}), color=GRAY)\n'
                f'# Header bar\n'
                f'hdr = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches({content_left}), Inches({content_top}), Inches({content_width}), Inches(0.7))\n'
                f'hdr.fill.solid()\n'
                f'hdr.fill.fore_color.rgb = ACCENT\n'
                f'hdr.line.fill.background()\n'
                f'hdr.text_frame.margin_left = Inches(0.3)\n'
                f'set_text(hdr.text_frame, "Section Title", Pt({header_pt}), bold=True, color=WHITE)\n'
                f'# Body\n'
                f'tf = slide.shapes.add_textbox(Inches({content_left}), Inches({round(content_top + 0.85, 2)}), Inches({content_width}), Inches({round(content_bottom - content_top - 1.0, 2)})).text_frame\n'
                f'tf.word_wrap = True\n'
                f'set_text(tf, "Content here...", Pt({body_pt}), color=DK2)\n\n'
                f'# --- Closing slide (1 text box) ---\n'
                f'slide = prs.slides.add_slide(title_layout)\n'
                f'strip_placeholders(slide)\n'
                f'tf = slide.shapes.add_textbox(Inches({title_box["x"]}), Inches({round(title_box["y"] + 0.6, 2)}), Inches({title_box["w"]}), Inches(1.5)).text_frame\n'
                f'set_text(tf, "Thank You", Pt({title_pt}), bold=True)\n\n'
                f'prs.save("{output_path}")\n'
                f'```\n\n'
                f'**Rules:**\n'
                f'- Always open with `Presentation("{template_path}")` — never `Presentation()`\n'
                f'- Template has ZERO slides — just add new ones\n'
                f'- EVERY slide: call strip_placeholders() immediately after adding\n'
                f'- EVERY text run: set `font.name = "{font_name}"` explicitly\n'
                f'- Title slide: EXACTLY 3 text boxes (title/subtitle/date) — no italic taglines, no "Prepared by", no extra lines\n'
                f'- Closing slide: EXACTLY 1 text box (e.g. "Thank You") — nothing else\n'
                f'- Title/closing slides: use `title_layout`, text at x={title_box["x"]}"\n'
                f'- Content slides: use `content_layout` — master auto-provides small logo + footer\n'
                f'- Content area: y={content_top}" to {content_bottom}" (logo ends at {logo_h}")\n'
                f'- If content is too dense, split across multiple slides\n'
                f'- Write ONE complete Python script and execute with `python3 script.py`\n'
                f'- After saving, run: `python3 /app/skills/pptx/scripts/office/validate.py {output_path}`\n'
                f'- Fix any validation errors before declaring DONE\n'
                f'- Output to: {output_path}\n'
            )
        else:
            template_section = (
                f'## Approach — NO TEMPLATE (python-pptx from scratch)\n\n'
                f'Use **python-pptx** to create a new presentation from scratch.\n\n'
                f'```python\n'
                f'from pptx import Presentation\n'
                f'from pptx.util import Inches, Pt\n'
                f'from pptx.dml.color import RGBColor\n'
                f'from pptx.enum.shapes import MSO_SHAPE\n'
                f'from pptx.enum.text import PP_ALIGN\n\n'
                f'prs = Presentation()\n'
                f'prs.slide_width = Inches(13.33)\n'
                f'prs.slide_height = Inches(7.5)\n'
                f'blank_layout = prs.slide_layouts[6]  # blank layout\n'
                f'```\n\n'
                f'Use `prs.slide_layouts[6]` (blank) for all slides — add shapes/text manually.\n'
                f'Set backgrounds with `slide.background.fill.solid()` and `fill.fore_color.rgb = ...`\n'
                f'Output to: {output_path}\n'
            )

        env_section = (
            f'## Environment\n\n'
            f'`python-pptx` is pre-installed. Run scripts with `python3 script.py`.\n'
            f'`react-icons`, `react`, `react-dom`, and `sharp` are available via Node.js for icon rendering.\n'
            f'Before running node commands: `export NODE_PATH=/usr/local/lib/node_modules`\n'
            f'Do NOT spawn sub-agents or use the Agent tool.\n'
            f'Do NOT read slide images or convert to PDF/JPEG — visual QA is handled externally.\n\n'
        )

        prompt = (
            f'{skill_prompt}\n\n'
            f'---\n\n'
            f'{template_section}\n'
            f'---\n\n'
            f'{env_section}'
            f'---\n\n'
            f'## Task\n\n'
            f'Generate a PowerPoint presentation from the following content.\n'
            f'\nTitle: {body.title}\n\n'
            f'Content:\n{body.content}\n\n'
            f'Write the complete file to: {output_path}\n'
            f'Print `DONE: {output_path}` when complete.\n'
        )

        stderr_lines: list[str] = []

        def _capture_stderr(line: str) -> None:
            stderr_lines.append(line)
            logger.warning('[CLAUDE-SERVICE] pptx-agent stderr: %s', line.rstrip())

        agent_environment = {
            **env,
            'NODE_PATH': '/usr/local/lib/node_modules',
        }

        options = ClaudeAgentOptions(
            permission_mode='bypassPermissions',
            tools=['Bash', 'Write'],
            allowed_tools=['Bash', 'Write'],
            model=body.model_id,
            max_turns=25,
            env=agent_environment,
            cwd=str(output_dir),
            stderr=_capture_stderr,
            max_buffer_size=10 * 1024 * 1024,
        )

        if body.job_id:
            await post_event(body.job_id, {'type': 'subagent_progress', 'label': 'Generating presentation...'})

        last_tool_action = ''
        _cancel_watcher = _start_cancel_watch(body.job_id) if body.job_id else None
        try:
            async for message in query(prompt=prompt, options=options):
                if isinstance(message, AssistantMessage):
                    for block in message.content:
                        if isinstance(block, ToolUseBlock):
                            gen_step += 1
                            raw_input = block.input if isinstance(block.input, dict) else {}
                            last_tool_action = f'{block.name}({list(raw_input.keys())})'
                            logger.info('[CLAUDE-SERVICE] pptx step %d: %s', gen_step, last_tool_action)
                            cmd = raw_input.get('command', '')
                            # Capture from Write tool (agent writes script to file)
                            if block.name == 'Write':
                                content = raw_input.get('content', '')
                                file_path = raw_input.get('file_path', '')
                                if content and file_path.endswith('.py') and 'from pptx' in content:
                                    source_script = content
                            # Capture from Bash if it's an inline script
                            elif cmd and (_is_pptx_script(cmd) or output_path.stem in cmd):
                                if 'from pptx' in cmd and len(cmd) > 200:
                                    source_script = cmd
                            if body.job_id:
                                pending_tool_calls[block.id] = {'tool': block.name, 'step': gen_step}
                                await post_event(body.job_id, {
                                    'type': 'subagent_tool_call',
                                    'tool': block.name,
                                    'args': _sanitize_tool_args(block.name, raw_input),
                                    'step': gen_step,
                                })
                        elif isinstance(block, ThinkingBlock) and body.job_id:
                            await post_event(body.job_id, {
                                'type': 'subagent_thinking',
                                'step': gen_step,
                                'length': len(block.thinking),
                            })
                elif isinstance(message, UserMessage) and body.job_id:
                    if isinstance(message.content, list):
                        for block in message.content:
                            if isinstance(block, ToolResultBlock):
                                call_info = pending_tool_calls.pop(block.tool_use_id, None)
                                await post_event(body.job_id, {
                                    'type': 'subagent_tool_result',
                                    'tool': call_info['tool'] if call_info else 'unknown',
                                    'step': call_info['step'] if call_info else gen_step,
                                    'contentLength': len(block.content) if isinstance(block.content, str) else 0,
                                    'isError': block.is_error or False,
                                })
                elif isinstance(message, ResultMessage) and body.job_id:
                    await post_event(body.job_id, {
                        'type': 'subagent_complete',
                        'label': 'Presentation generation complete',
                        'durationMs': message.duration_ms,
                    })
        except asyncio.CancelledError:
            logger.info('[CLAUDE-SERVICE] job cancelled, stopping create-pptx')
        except Exception as exc:
            logger.warning(
                '[CLAUDE-SERVICE] pptx agent stopped after %d steps. last_action=%s: %s',
                gen_step, last_tool_action, exc,
            )
        finally:
            if _cancel_watcher:
                _cancel_watcher.cancel()

        if not output_path.exists():
            raise HTTPException(status_code=422, detail='Presentation generation failed — agent did not produce a file')

        content = base64.b64encode(output_path.read_bytes()).decode('ascii')
        logger.info('[CLAUDE-SERVICE] create-pptx complete')

        return {
            'artifacts': {
                'content': content,
                'extension': '.pptx',
                'encoding': 'base64',
            },
            'source_script': source_script,
        }
    finally:
        shutil.rmtree(output_dir, ignore_errors=True)


class EditPptxRequest(BaseModel):
    model_id: str
    internal_model_id: str = ''
    user_id: str
    chat_message_id: str | None = None
    user_group_id: str | None = None
    source_script: str
    changes: str
    job_id: str | None = None


@router.post('/edit-pptx')
async def edit_pptx(request: Request, body: EditPptxRequest):
    verify_auth(request)

    output_dir = Path(tempfile.mkdtemp(prefix='pptx-edit-'))
    output_path = output_dir / 'presentation.pptx'

    try:
        logger.info('[CLAUDE-SERVICE] edit-pptx changes=%r source_script_len=%d', body.changes[:120], len(body.source_script))

        skill_prompt = load_skill('pptx')

        if body.job_id:
            await post_event(body.job_id, {'type': 'subagent_progress', 'label': 'Applying edits to presentation...'})

        # Re-resolve template if the script references one (the original temp dir was cleaned up)
        source_script = body.source_script
        if 'Presentation(' in source_script and 'template.pptx' in source_script:
            env = agent_env(body.user_id, body.chat_message_id, body.user_group_id)
            template_bytes = await _resolve_pptx_template(body.user_id, env)
            if template_bytes:
                template_path = output_dir / 'template.pptx'
                template_path.write_bytes(template_bytes)
                _fix_potx_content_type(template_path)
                _strip_template_slides(template_path)
                # Rewrite the old temp path to the new one
                source_script = re.sub(
                    r'Presentation\(["\'][^"\']*template\.pptx["\']\)',
                    f'Presentation("{template_path}")',
                    source_script,
                )

        is_legacy = 'slide_factory' in source_script

        if is_legacy:
            prompt = (
                f'{skill_prompt}\n\n---\n\n'
                f'## Task\n\n'
                f'The previous version of this presentation used an outdated generation method.\n'
                f'Regenerate it from scratch using python-pptx, incorporating these changes:\n\n'
                f'Requested changes:\n{body.changes}\n\n'
                f'The previous script (for content reference only — do NOT reuse this approach):\n'
                f'```\n{source_script}\n```\n\n'
                f'Write the complete file to: {output_path}\n'
                f'Print `DONE: {output_path}` when complete.\n'
            )
        else:
            prompt = (
                f'{skill_prompt}\n\n---\n\n'
                f'## Task\n\n'
                f'Edit the existing PowerPoint presentation by applying the requested changes.\n\n'
                f'Requested changes:\n{body.changes}\n\n'
                f'Here is the python-pptx script that generated the current presentation.\n'
                f'Modify it to incorporate the changes, then execute with `python3`.\n'
                f'The template file is already on disk at the path referenced in the script.\n\n'
                f'```python\n{source_script}\n```\n\n'
                f'Do NOT explore the filesystem, read XML, or open the .pptx file.\n'
                f'Just edit the script above, save it, run it with python3, and validate.\n'
                f'Complete the edit in under 10 tool calls.\n'
                f'Write the complete file to: {output_path}\n'
                f'Print `DONE: {output_path}` when complete.\n'
            )

        edit_env = {
            **agent_env(body.user_id, body.chat_message_id, body.user_group_id),
            'NODE_PATH': '/usr/local/lib/node_modules',
        }

        stderr_lines: list[str] = []

        def _edit_stderr(line: str) -> None:
            stderr_lines.append(line)
            logger.warning('[CLAUDE-SERVICE] edit-agent stderr: %s', line.rstrip())

        options = ClaudeAgentOptions(
            permission_mode='bypassPermissions',
            tools=['Bash', 'Write'],
            allowed_tools=['Bash', 'Write'],
            model=body.model_id,
            max_turns=25,
            env=edit_env,
            cwd=str(output_dir),
            stderr=_edit_stderr,
            max_buffer_size=10 * 1024 * 1024,
        )

        new_source_script = ''
        gen_step, pending_tool_calls = 0, {}
        _cancel_watcher = _start_cancel_watch(body.job_id) if body.job_id else None
        try:
            async for message in query(prompt=prompt, options=options):
                if isinstance(message, AssistantMessage):
                    for block in message.content:
                        if isinstance(block, ToolUseBlock):
                            gen_step += 1
                            raw_input = block.input if isinstance(block.input, dict) else {}
                            cmd = raw_input.get('command', '')
                            if block.name == 'Write':
                                content = raw_input.get('content', '')
                                file_path = raw_input.get('file_path', '')
                                if content and file_path.endswith('.py') and 'from pptx' in content:
                                    new_source_script = content
                            elif cmd and _is_pptx_script(cmd):
                                new_source_script = cmd
                            if body.job_id:
                                pending_tool_calls[block.id] = {'tool': block.name, 'step': gen_step}
                                await post_event(body.job_id, {
                                    'type': 'subagent_tool_call',
                                    'tool': block.name,
                                    'args': _sanitize_tool_args(block.name, raw_input),
                                    'step': gen_step,
                                })
                        elif isinstance(block, ThinkingBlock) and body.job_id:
                            await post_event(body.job_id, {
                                'type': 'subagent_thinking',
                                'step': gen_step,
                                'length': len(block.thinking),
                            })
                elif isinstance(message, UserMessage) and body.job_id:
                    if isinstance(message.content, list):
                        for block in message.content:
                            if isinstance(block, ToolResultBlock):
                                call_info = pending_tool_calls.pop(block.tool_use_id, None)
                                await post_event(body.job_id, {
                                    'type': 'subagent_tool_result',
                                    'tool': call_info['tool'] if call_info else 'unknown',
                                    'step': call_info['step'] if call_info else gen_step,
                                    'contentLength': len(block.content) if isinstance(block.content, str) else 0,
                                    'isError': block.is_error or False,
                                })
                elif isinstance(message, ResultMessage) and body.job_id:
                    await post_event(body.job_id, {
                        'type': 'subagent_complete',
                        'label': 'Presentation edit complete',
                        'durationMs': message.duration_ms,
                    })
        except asyncio.CancelledError:
            logger.info('[CLAUDE-SERVICE] job cancelled, stopping edit-pptx')
        except Exception as exc:
            logger.warning(
                '[CLAUDE-SERVICE] edit-pptx agent stopped after %d steps: %s',
                gen_step, exc,
            )
        finally:
            if _cancel_watcher:
                _cancel_watcher.cancel()

        if not output_path.exists():
            raise HTTPException(status_code=422, detail='Presentation edit failed — agent did not produce a file')

        content = base64.b64encode(output_path.read_bytes()).decode('ascii')
        logger.info('[CLAUDE-SERVICE] edit-pptx complete')

        final_source_script = new_source_script or body.source_script
        diff_stat = None
        try:
            diff_stat = _line_diff_stat(body.source_script, final_source_script)
        except Exception:
            pass

        return {
            'artifacts': {
                'content': content,
                'extension': '.pptx',
                'encoding': 'base64',
            },
            'source_script': final_source_script,
            'diffStat': diff_stat,
        }
    finally:
        shutil.rmtree(output_dir, ignore_errors=True)


class DocxPreviewRequest(BaseModel):
    binaryContent: str
    artifactId: str


@router.post('/docx-preview')
async def docx_preview(request: Request, body: DocxPreviewRequest):
    verify_auth(request)

    tmp_dir = Path(tempfile.mkdtemp(prefix='docx-preview-'))
    try:
        docx_path = tmp_dir / 'document.docx'
        docx_path.write_bytes(base64.b64decode(body.binaryContent))

        lo_proc = await asyncio.create_subprocess_exec(
            'libreoffice', '--headless', '--convert-to', 'pdf',
            '--outdir', str(tmp_dir), str(docx_path),
            stdout=asyncio.subprocess.PIPE,
            stderr=asyncio.subprocess.PIPE,
        )
        _, lo_stderr = await lo_proc.communicate()
        if lo_proc.returncode != 0:
            raise HTTPException(
                status_code=422,
                detail=f'LibreOffice conversion failed: {lo_stderr.decode(errors="replace")}',
            )

        pdf_path = tmp_dir / 'document.pdf'
        if not pdf_path.exists():
            raise HTTPException(status_code=422, detail='PDF not produced by LibreOffice')

        page_prefix = str(tmp_dir / 'page')
        pp_proc = await asyncio.create_subprocess_exec(
            'pdftoppm', '-png', '-r', '150', str(pdf_path), page_prefix,
            stdout=asyncio.subprocess.PIPE,
            stderr=asyncio.subprocess.PIPE,
        )
        _, pp_stderr = await pp_proc.communicate()
        if pp_proc.returncode != 0:
            raise HTTPException(
                status_code=422,
                detail=f'pdftoppm failed: {pp_stderr.decode(errors="replace")}',
            )

        page_files = sorted(tmp_dir.glob('page-*.png'))
        pages = [base64.b64encode(f.read_bytes()).decode('ascii') for f in page_files]

        logger.info('[CLAUDE-SERVICE] docx-preview artifactId=%s pages=%d', body.artifactId, len(pages))
        return {'pages': pages, 'count': len(pages)}
    finally:
        shutil.rmtree(tmp_dir, ignore_errors=True)


class PptxPreviewRequest(BaseModel):
    binaryContent: str
    artifactId: str


@router.post('/pptx-preview')
async def pptx_preview(request: Request, body: PptxPreviewRequest):
    verify_auth(request)

    tmp_dir = Path(tempfile.mkdtemp(prefix='pptx-preview-'))
    try:
        pptx_path = tmp_dir / 'presentation.pptx'
        pptx_path.write_bytes(base64.b64decode(body.binaryContent))

        lo_proc = await asyncio.create_subprocess_exec(
            'libreoffice', '--headless', '--convert-to', 'pdf',
            '--outdir', str(tmp_dir), str(pptx_path),
            stdout=asyncio.subprocess.PIPE,
            stderr=asyncio.subprocess.PIPE,
        )
        _, lo_stderr = await lo_proc.communicate()
        if lo_proc.returncode != 0:
            raise HTTPException(
                status_code=422,
                detail=f'LibreOffice conversion failed: {lo_stderr.decode(errors="replace")}',
            )

        pdf_path = tmp_dir / 'presentation.pdf'
        if not pdf_path.exists():
            raise HTTPException(status_code=422, detail='PDF not produced by LibreOffice')

        slide_prefix = str(tmp_dir / 'slide')
        pp_proc = await asyncio.create_subprocess_exec(
            'pdftoppm', '-png', '-r', '120', str(pdf_path), slide_prefix,
            stdout=asyncio.subprocess.PIPE,
            stderr=asyncio.subprocess.PIPE,
        )
        _, pp_stderr = await pp_proc.communicate()
        if pp_proc.returncode != 0:
            raise HTTPException(
                status_code=422,
                detail=f'pdftoppm failed: {pp_stderr.decode(errors="replace")}',
            )

        png_files = sorted(tmp_dir.glob('slide-*.png'))
        slides = [base64.b64encode(f.read_bytes()).decode('ascii') for f in png_files]

        logger.info('[CLAUDE-SERVICE] pptx-preview artifactId=%s slides=%d', body.artifactId, len(slides))
        return {'slides': slides, 'count': len(slides)}
    finally:
        shutil.rmtree(tmp_dir, ignore_errors=True)


@router.post('/create-xlsx')
async def create_xlsx(request: Request, body: CreateXlsxRequest):
    verify_auth(request)

    output_dir = Path(tempfile.mkdtemp(prefix='xlsx-'))
    output_path = output_dir / 'spreadsheet.xlsx'

    try:
        logger.info('[CLAUDE-SERVICE] create-xlsx title=%r', body.title[:60])

        skill_prompt = load_skill('xlsx')
        prompt = (
            f'{skill_prompt}\n\n'
            f'---\n\n'
            f'## Task\n\n'
            f'Create a spreadsheet from the following content.\n\n'
            f'Title: {body.title}\n\n'
            f'Content:\n{body.content}\n\n'
            f'Write the complete file to: {output_path}\n'
            f'Print `DONE: {output_path}` when complete.\n'
        )

        options = ClaudeAgentOptions(
            tools=['Bash'],
            allowed_tools=['Bash'],
            model=body.model_id,
            env=agent_env(body.user_id, body.chat_message_id, body.user_group_id),
        )

        if body.job_id:
            await post_event(body.job_id, {'type': 'subagent_progress', 'label': 'Building spreadsheet...'})

        source_script: str | None = None
        gen_step = 0
        pending_tool_calls: dict[str, dict] = {}
        _cancel_watcher = _start_cancel_watch(body.job_id) if body.job_id else None
        try:
            async for message in query(prompt=prompt, options=options):
                if isinstance(message, AssistantMessage):
                    for block in message.content:
                        if isinstance(block, ToolUseBlock):
                            gen_step += 1
                            raw_input = block.input if isinstance(block.input, dict) else {}
                            pending_tool_calls[block.id] = {'tool': block.name, 'step': gen_step}
                            if block.name == 'Bash' and isinstance(raw_input, dict):
                                cmd = raw_input.get('command', '')
                                if 'openpyxl' in cmd or 'workbook' in cmd.lower() or '.xlsx' in cmd:
                                    source_script = cmd
                            if body.job_id:
                                await post_event(body.job_id, {
                                    'type': 'subagent_tool_call',
                                    'tool': block.name,
                                    'args': _sanitize_tool_args(block.name, raw_input),
                                    'step': gen_step,
                                })
                        elif isinstance(block, ThinkingBlock) and body.job_id:
                            await post_event(body.job_id, {
                                'type': 'subagent_thinking',
                                'step': gen_step,
                                'length': len(block.thinking),
                            })
                elif isinstance(message, UserMessage) and body.job_id:
                    if isinstance(message.content, list):
                        for block in message.content:
                            if isinstance(block, ToolResultBlock):
                                call_info = pending_tool_calls.pop(block.tool_use_id, None)
                                content_len = len(block.content) if isinstance(block.content, str) else 0
                                await post_event(body.job_id, {
                                    'type': 'subagent_tool_result',
                                    'tool': call_info['tool'] if call_info else 'unknown',
                                    'step': call_info['step'] if call_info else gen_step,
                                    'contentLength': content_len,
                                    'isError': block.is_error or False,
                                })
                elif isinstance(message, ResultMessage) and body.job_id:
                    await post_event(body.job_id, {
                        'type': 'subagent_complete',
                        'label': 'Spreadsheet generation complete',
                        'durationMs': message.duration_ms,
                    })
        except asyncio.CancelledError:
            logger.info('[CLAUDE-SERVICE] job cancelled, stopping create-xlsx')
        finally:
            if _cancel_watcher:
                _cancel_watcher.cancel()

        if not output_path.exists():
            raise HTTPException(status_code=422, detail='Spreadsheet generation failed — agent did not produce a file')

        raw = output_path.read_bytes()
        content = base64.b64encode(raw).decode('ascii')

        return {
            'artifacts': {
                'content': content,
                'extension': '.xlsx',
                'encoding': 'base64',
            },
            'source_script': source_script,
        }
    finally:
        shutil.rmtree(output_dir, ignore_errors=True)


@router.post('/edit-xlsx')
async def edit_xlsx(request: Request, body: EditXlsxRequest):
    verify_auth(request)

    output_dir = Path(tempfile.mkdtemp(prefix='xlsx-edit-'))
    raw_suffix = Path(body.filename).suffix.lower() if body.filename else '.xlsx'
    input_suffix = raw_suffix if raw_suffix in ('.xlsx', '.csv') else '.xlsx'
    is_csv_input = input_suffix == '.csv'
    extension = '.csv' if is_csv_input else '.xlsx'
    output_filename = f'spreadsheet{extension}'
    output_path = output_dir / output_filename

    try:
        logger.info('[CLAUDE-SERVICE] edit-xlsx title=%r changes=%r', body.title[:60], body.changes[:120])

        skill_prompt = load_skill('xlsx')

        if body.data_b64:
            # Path B — uploaded file: write raw data to disk and have the subagent manipulate it directly
            input_path = output_dir / f'input{input_suffix}'
            input_path.write_bytes(base64.b64decode(body.data_b64))
            output_format_note = 'Save the result as a CSV file (.csv).' if is_csv_input else 'Save the result as an Excel file (.xlsx).'
            prompt = (
                f'{skill_prompt}\n\n'
                f'---\n\n'
                f'## Task\n\n'
                f'Edit the uploaded spreadsheet by applying the requested changes.\n\n'
                f'Input file: {input_path}\n'
                f'Requested changes:\n{body.changes}\n\n'
                f'Load the input file, apply the changes, then {output_format_note}\n'
                f'Write the complete file to: {output_path}\n'
                f'Print `DONE: {output_path}` when complete.\n'
            )
        else:
            # Path A — generated artifact: rewrite the source script
            prompt = (
                f'{skill_prompt}\n\n'
                f'---\n\n'
                f'## Task\n\n'
                f'Edit the existing spreadsheet by applying the requested changes to the script below.\n\n'
                f'Title: {body.title}\n\n'
                f'Requested changes:\n{body.changes}\n\n'
                f'Existing script:\n```python\n{body.source_script}\n```\n\n'
                f'Rewrite the script incorporating the changes, then execute it to produce the updated file.\n'
                f'Write the complete file to: {output_path}\n'
                f'Print `DONE: {output_path}` when complete.\n'
            )

        options = ClaudeAgentOptions(
            tools=['Bash'],
            allowed_tools=['Bash'],
            model=body.model_id,
            env=agent_env(body.user_id, body.chat_message_id, body.user_group_id),
        )

        if body.job_id:
            await post_event(body.job_id, {'type': 'subagent_progress', 'label': 'Applying edits to spreadsheet...'})

        source_script: str | None = None
        gen_step = 0
        pending_tool_calls: dict[str, dict] = {}
        _cancel_watcher = _start_cancel_watch(body.job_id) if body.job_id else None
        try:
            async for message in query(prompt=prompt, options=options):
                if isinstance(message, AssistantMessage):
                    for block in message.content:
                        if isinstance(block, ToolUseBlock):
                            gen_step += 1
                            raw_input = block.input if isinstance(block.input, dict) else {}
                            pending_tool_calls[block.id] = {'tool': block.name, 'step': gen_step}
                            if block.name == 'Bash' and isinstance(raw_input, dict):
                                cmd = raw_input.get('command', '')
                                if 'openpyxl' in cmd or 'workbook' in cmd.lower() or '.xlsx' in cmd:
                                    source_script = cmd
                            if body.job_id:
                                await post_event(body.job_id, {
                                    'type': 'subagent_tool_call',
                                    'tool': block.name,
                                    'args': _sanitize_tool_args(block.name, raw_input),
                                    'step': gen_step,
                                })
                        elif isinstance(block, ThinkingBlock) and body.job_id:
                            await post_event(body.job_id, {
                                'type': 'subagent_thinking',
                                'step': gen_step,
                                'length': len(block.thinking),
                            })
                elif isinstance(message, UserMessage) and body.job_id:
                    if isinstance(message.content, list):
                        for block in message.content:
                            if isinstance(block, ToolResultBlock):
                                call_info = pending_tool_calls.pop(block.tool_use_id, None)
                                content_len = len(block.content) if isinstance(block.content, str) else 0
                                await post_event(body.job_id, {
                                    'type': 'subagent_tool_result',
                                    'tool': call_info['tool'] if call_info else 'unknown',
                                    'step': call_info['step'] if call_info else gen_step,
                                    'contentLength': content_len,
                                    'isError': block.is_error or False,
                                })
                elif isinstance(message, ResultMessage) and body.job_id:
                    await post_event(body.job_id, {
                        'type': 'subagent_complete',
                        'label': 'Spreadsheet edit complete',
                        'durationMs': message.duration_ms,
                    })
        except asyncio.CancelledError:
            logger.info('[CLAUDE-SERVICE] job cancelled, stopping edit-xlsx')
        finally:
            if _cancel_watcher:
                _cancel_watcher.cancel()

        if not output_path.exists():
            raise HTTPException(status_code=422, detail='Spreadsheet edit failed — agent did not produce a file')

        if is_csv_input:
            content = output_path.read_text(encoding='utf-8')
            encoding = 'utf-8'
        else:
            content = base64.b64encode(output_path.read_bytes()).decode('ascii')
            encoding = 'base64'

        diff_stat = None
        try:
            if body.data_b64:
                before_text = input_path.read_text(encoding='utf-8') if is_csv_input else _xlsx_text(input_path)
                after_text = output_path.read_text(encoding='utf-8') if is_csv_input else _xlsx_text(output_path)
            else:
                before_text = body.source_script
                after_text = source_script or body.source_script
            diff_stat = _line_diff_stat(before_text, after_text)
        except Exception:
            pass

        return {
            'artifacts': {
                'content': content,
                'extension': extension,
                'encoding': encoding,
            },
            'source_script': source_script,
            'diffStat': diff_stat,
        }
    finally:
        shutil.rmtree(output_dir, ignore_errors=True)


@router.post('/create-mp4')
async def create_mp4(request: Request, body: CreateMp4Request):
    verify_auth(request)

    job_id = body.job_id or os.urandom(8).hex()
    output_dir = Path(tempfile.mkdtemp(prefix='mp4-'))
    output_path = output_dir / 'video.mp4'

    try:
        logger.info('[CLAUDE-SERVICE] create-mp4 title=%r', body.title[:60])

        result = await _run_agent_skill(
            skill_name='video',
            output_type='video',
            output_path=output_path,
            model_id=body.model_id,
            user_id=body.user_id,
            chat_message_id=body.chat_message_id,
            user_group_id=body.user_group_id,
            title=body.title,
            content=body.content,
            style=body.style,
            feedback=body.feedback,
            job_id=body.job_id,
        )

        if not output_path.exists():
            raise HTTPException(status_code=422, detail='Video generation failed')

        content = output_path.read_text(encoding='utf-8')

        return {'artifacts': {
            'content': content,
            'extension': '.mp4',
            'encoding': 'utf-8',
        }}
    finally:
        shutil.rmtree(output_dir, ignore_errors=True)


async def _run_agent_skill(
    skill_name: str,
    output_type: str,
    output_path: Path,
    model_id: str,
    user_id: str,
    chat_message_id: str | None,
    user_group_id: str | None,
    title: str,
    content: str,
    style: str,
    feedback: str,
    job_id: str | None,
) -> None:
    if job_id:
        await post_event(job_id, {
            'type': 'subagent_progress',
            'label': f'Loading {output_type} skill...',
        })

    skill_prompt = load_skill(skill_name)
    effective_style = style if style and style != 'default' else 'corporate-minimal'
    style_hint = f'\nDesign style preference: {effective_style}\n'

    feedback_hint = ''
    if feedback:
        feedback_hint = f'\nFeedback on previous version (apply these changes):\n{feedback}\n'

    prompt = (
        f'{skill_prompt}\n\n'
        f'---\n\n'
        f'## Task\n\n'
        f'Generate a {output_type} from the following content.\n'
        f'{style_hint}'
        f'{feedback_hint}'
        f'\nTitle: {title}\n\n'
        f'Content:\n{content}\n\n'
        f'Write the complete file to: {output_path}\n'
    )

    options = ClaudeAgentOptions(
        tools=['Bash'],
        allowed_tools=['Bash'],
        model=model_id,
        env=agent_env(user_id, chat_message_id, user_group_id),
    )

    if job_id:
        await post_event(job_id, {
            'type': 'subagent_progress',
            'label': f'Generating {output_type}...',
        })

    gen_step = 0
    pending_tool_calls: dict[str, dict] = {}
    last_tool_errors: list[str] = []
    _cancel_watcher = _start_cancel_watch(job_id) if job_id else None
    try:
        async for message in query(prompt=prompt, options=options):
            if not job_id:
                continue
            if isinstance(message, AssistantMessage):
                for block in message.content:
                    if isinstance(block, ToolUseBlock):
                        gen_step += 1
                        raw_input = block.input if isinstance(block.input, dict) else {}
                        pending_tool_calls[block.id] = {'tool': block.name, 'step': gen_step}
                        await post_event(job_id, {
                            'type': 'subagent_tool_call',
                            'tool': block.name,
                            'args': _sanitize_tool_args(block.name, raw_input),
                            'step': gen_step,
                        })
                    elif isinstance(block, ThinkingBlock):
                        await post_event(job_id, {
                            'type': 'subagent_thinking',
                            'step': gen_step,
                            'length': len(block.thinking),
                        })
            elif isinstance(message, UserMessage):
                if isinstance(message.content, list):
                    for block in message.content:
                        if isinstance(block, ToolResultBlock):
                            call_info = pending_tool_calls.pop(block.tool_use_id, None)
                            tool_name = call_info['tool'] if call_info else 'unknown'
                            content_str = block.content if isinstance(block.content, str) else str(block.content)
                            content_len = len(content_str)
                            if block.is_error:
                                snippet = content_str[:500]
                                logger.error('[CLAUDE-SERVICE] %s tool error on step %d: %s', tool_name, call_info['step'] if call_info else gen_step, snippet)
                                last_tool_errors.append(f'{tool_name}: {snippet}')
                            await post_event(job_id, {
                                'type': 'subagent_tool_result',
                                'tool': tool_name,
                                'step': call_info['step'] if call_info else gen_step,
                                'contentLength': content_len,
                                'isError': block.is_error or False,
                            })
            elif isinstance(message, ResultMessage):
                await post_event(job_id, {
                    'type': 'subagent_complete',
                    'label': f'{output_type.capitalize()} generation complete',
                    'durationMs': message.duration_ms,
                })
    except asyncio.CancelledError:
        logger.info('[CLAUDE-SERVICE] job cancelled, stopping %s generation', output_type)
    except Exception as exc:
        tool_error_ctx = '; '.join(last_tool_errors[-3:]) if last_tool_errors else 'no tool errors captured'
        logger.error(
            '[CLAUDE-SERVICE] %s agent failed (step %d): %s | recent tool errors: %s',
            skill_name, gen_step, exc, tool_error_ctx,
        )
        raise
    finally:
        if _cancel_watcher:
            _cancel_watcher.cancel()


def _line_diff_stat(before: str, after: str) -> dict:
    diff = list(difflib.ndiff(before.splitlines(), after.splitlines()))
    return {
        'added': sum(1 for l in diff if l.startswith('+ ')),
        'removed': sum(1 for l in diff if l.startswith('- ')),
    }


def _xlsx_text(path: Path) -> str:
    from openpyxl import load_workbook
    wb = load_workbook(str(path), data_only=True)
    lines = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            lines.append('\t'.join('' if c is None else str(c) for c in row))
    return '\n'.join(lines)


def _sanitize_tool_args(tool_name: str, raw_input: dict) -> dict:
    if tool_name == 'Bash':
        return {'command': raw_input.get('command', '')}
    if tool_name == 'Write':
        return {'file_path': raw_input.get('file_path', '')}
    if tool_name == 'Read':
        return {'file_path': raw_input.get('file_path', '')}
    if tool_name == 'Edit':
        return {'file_path': raw_input.get('file_path', '')}
    return {k: str(v) for k, v in list(raw_input.items())[:5]}
